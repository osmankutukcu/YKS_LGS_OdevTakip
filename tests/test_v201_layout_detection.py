"""V20.1 structural layout recognition. All tests are read-only and independent of Qt."""
import csv
from pathlib import Path
import pytest
from openpyxl import Workbook
from services.question_evidence_v19 import parse_analysis_file, ReportFormatError, prepare
from services.question_evidence_v19.layout_detector_v201 import match_header
from tests.test_v19_question_topic_evidence import db


def write_csv(tmp_path, name, lines):
    p=tmp_path/name
    p.write_text(lines,encoding='utf-8-sig')
    return p


def workbook(tmp_path, data, name='karne.xlsx'):
    wb=Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in data:
        sh=wb.create_sheet(sheet_name)
        for row in rows:sh.append(row)
    p=tmp_path/name
    wb.save(p)
    return p


def test_exact_aliases_recognize_variant_names():
    found=match_header(['Branş','Soru Num','D/Y/B','Konu Başlığı','Kazanım No'])
    assert found and all(x in found.columns for x in ('ders','soru_no','sonuc','konu','kazanim_kodu'))


def test_only_net_summary_never_structural_question():
    assert match_header(['Branş','Doğru','Yanlış','Net']) is None


def test_duplicate_topic_alias_fails_closed():
    with pytest.raises(ValueError,match='Tekrarlanan'):
        match_header(['Ders','Soru No','Sonuç','Konu','Konu Adı'])


def test_multi_sheet_info_plus_exam_data(tmp_path):
    p=workbook(tmp_path,[('Kapak',[['Sınav ve kurum özeti'],['Öğrenci','Not']]),
                      ('Fizik',[['Öğrenci ID','Deneme ID','Branş','Soru Num','D/Y/B','Konu Başlığı'],
                                [1,9,'Fizik',3,'Yanlış','Elektrostatik']]),
                      ('Matematik',[['Öğrenci ID','Deneme ID','Branş','Soru Num','D/Y/B','Konu Başlığı'],
                                   [1,9,'Matematik',7,'Boş','Fonksiyonlar']])])
    r=parse_analysis_file(p,student_id=1,exam_id=9)
    assert len(r.rows)==2 and r.scanned_sheets==3
    assert 'Çok sayfalı' in r.detected_layout
    assert any('Kapak' in w for w in r.warnings)
    assert r.student_id_in_source and r.exam_id_in_source


def test_conflicting_identity_in_secondary_sheet_blocks_all(tmp_path):
    p=workbook(tmp_path,[('A',[['Ogrenci ID','Deneme ID','Ders','Soru No','Sonuc','Konu'],[1,9,'Fizik',3,'Y','Dalga']]),
                       ('B',[['Ogrenci ID','Deneme ID','Ders','Soru No','Sonuc','Konu'],[2,9,'Fizik',4,'Y','Dalga']])])
    with pytest.raises(ReportFormatError,match='B.*uyuşmuyor'):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_repeated_question_across_sheets_not_added_twice(tmp_path):
    p=workbook(tmp_path,[('A',[['Ders','Soru No','Sonuc','Konu'],['Fizik',3,'Y','Dalga']]),
                       ('B',[['Ders','Soru No','Sonuc','Konu'],['Fizik',3,'Y','Dalga']])])
    r=parse_analysis_file(p,student_id=1,exam_id=9)
    prev=prepare(db(),9,1,r.rows)
    assert prev.errors


def test_formula_rejected(tmp_path):
    p=workbook(tmp_path,[('Soru',[['Ders','Soru No','Sonuc','Konu'],['Fizik', '=1+2', 'Yanlış','Dalga']])])
    with pytest.raises(ReportFormatError,match='Formül'):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_only_overview_sheet_rejected(tmp_path):
    p=workbook(tmp_path,[('Ozet',[['Ders','Doğru','Yanlış','Net'],['Fizik',4,2,3.5]])])
    with pytest.raises(ReportFormatError,match='açık soru no'):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_non_question_data_does_not_create_student_evidence(tmp_path):
    p=write_csv(tmp_path,'summary.csv','Ders,Doğru,Yanlış,Boş,Net\nFizik,3,2,1,2.5\n')
    with pytest.raises(ReportFormatError):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_unverified_identity_still_flagged_per_sheet(tmp_path):
    p=workbook(tmp_path,[('Kimliksiz',[['Ders','Soru No','Sonuc','Konu'],['Fizik',3,'Yanlış','Dalga']])])
    r=parse_analysis_file(p,student_id=1,exam_id=9)
    assert not r.student_id_in_source and not r.exam_id_in_source
    assert any('öğrenci ID' in warning for warning in r.warnings)


def test_hidden_exam_sheet_not_read(tmp_path):
    p=workbook(tmp_path,[('Kapak',[['Sayfa başlığı']]),('Gizli',[['Ders','Soru No','Sonuc','Konu'],['Fizik',3,'Y','Dalga']])])
    from openpyxl import load_workbook
    wb=load_workbook(p)
    wb['Gizli'].sheet_state='hidden';wb.save(p)
    with pytest.raises(ReportFormatError,match='açık soru no'):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_pdf_photo_still_not_accepted(tmp_path):
    p=tmp_path/'scan.jpg';p.write_bytes(b'fake')
    with pytest.raises(ReportFormatError,match='Yalnızca'):
        parse_analysis_file(p,student_id=1,exam_id=9)


def test_partial_question_sheet_blocks_instead_of_silent_skip(tmp_path):
    p=workbook(tmp_path,[('Gecerli',[['Ders','Soru No','Sonuc','Konu'],['Fizik',3,'Y','Dalga']]),
                       ('Eksik',[['Ders','Soru No','Sonuc'],['Fizik',4,'Yanlış']])])
    with pytest.raises(ReportFormatError,match='Eksik.*eksik sütun'):
        parse_analysis_file(p,student_id=1,exam_id=9)
