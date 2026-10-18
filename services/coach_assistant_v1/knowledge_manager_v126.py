# -*- coding: utf-8 -*-
"""V12.6: reviewed, local, opt-in coaching knowledge extension.

User-supplied documents are untrusted drafts. Only explicitly reviewed entries
are searchable, and neither drafts nor active entries can execute instructions,
access student DB or send network requests. This is not scientific verification.
"""
from __future__ import annotations
import hashlib
import json
import os
import re
import tempfile
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

MAX_STORE_BYTES = 1_000_000
MAX_INPUT_BYTES = 6_000_000
MAX_CANDIDATES_PER_IMPORT = 50
MAX_ENTRIES = 300
CATEGORIES = {'coaching','homework','assessment','planning','learning_science','pedagogy','error_analysis','student_support','methodology'}


def fold(v):
    v = str(v or '').replace('ı', 'i').replace('İ', 'I')
    v = unicodedata.normalize('NFKD', v.casefold())
    return ' '.join(re.findall(r'[a-z0-9]+', ''.join(c for c in v if unicodedata.category(c) != 'Mn')))


def default_store_path():
    # per-OS-user configuration; must not write into application/project/student data
    if os.name == 'nt':
        root = Path(os.environ.get('APPDATA') or Path.home()/'AppData'/'Roaming')
    elif os.environ.get('XDG_CONFIG_HOME'):
        root = Path(os.environ['XDG_CONFIG_HOME'])
    elif os.sys.platform == 'darwin':
        root = Path.home()/'Library'/'Application Support'
    else:
        root = Path.home()/'.config'
    return root/'EduTrack'/'coach_knowledge_reviews_v126.json'


def _now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def _safe_str(s, cap=500):
    if not isinstance(s, str): raise ValueError('Metin bekleniyor')
    s = s.strip()
    if len(s)>cap: raise ValueError('Alan uzunluğu sınırını aşıyor')
    return s


def _extract_document(source:Path):
    if not source.is_file() or source.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError('Dosya bulunamadı veya çok büyük')
    suffix=source.suffix.lower()
    if suffix=='.docx':
        # Refuse compressed documents with excessive extracted content.
        from zipfile import ZipFile, BadZipFile
        try:
            with ZipFile(source) as archive:
                members=archive.infolist()
                if len(members)>500 or sum(m.file_size for m in members)>12_000_000:
                    raise ValueError('DOCX açıldığında boyut sınırını aşıyor')
        except BadZipFile:
            raise ValueError('Geçersiz DOCX arşivi') from None
        from docx import Document
        doc=Document(source)
        fragments=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # DOCX tables can hold useful lists and action notes
        for table in doc.tables[:12]:
            for row in table.rows[:100]:
                txt=' | '.join(c.text.strip().replace('\n','; ') for c in row.cells)
                if txt.strip(' |'): fragments.append(txt)
    elif suffix in ('.txt','.md'):
        fragments=[p.strip() for p in re.split(r'\n\s*\n',source.read_text('utf-8-sig')) if p.strip()]
    else:
        raise ValueError('Yalnızca DOCX, TXT ve MD desteklenir')
    out=[]
    for p in fragments:
        p=' '.join(p.split())
        if 25<=len(p)<=1800 and p not in out:
            out.append(p)
        if len(out)>=MAX_CANDIDATES_PER_IMPORT: break
    return out


class ReviewStore:
    def __init__(self,path=None):
        self.path=Path(path) if path is not None else default_store_path()

    def read(self):
        if not self.path.exists():
            return {'schema_version':'1.0','entries':[]}
        if self.path.stat().st_size>MAX_STORE_BYTES:
            raise ValueError('Bilgi kütüphanesi sınırı aşıldı')
        data=json.loads(self.path.read_text(encoding='utf-8'))
        if not isinstance(data,dict) or data.get('schema_version')!='1.0' or not isinstance(data.get('entries'),list) or len(data['entries'])>MAX_ENTRIES:
            raise ValueError('Bilgi kütüphanesi biçimi hatalı')
        seen=set()
        for e in data['entries']:
            if not isinstance(e,dict) or not isinstance(e.get('id'),str) or e['id'] in seen or e.get('status') not in ('draft','reviewed','retired'):
                raise ValueError('Onay kaydı geçersiz')
            if e.get('automatic_action_allowed') is not False:
                raise ValueError('Otomatik işlem içeren kayıt kabul edilmez')
            seen.add(e['id'])
        return data

    def _save(self,data):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if len(data.get('entries',[]))>MAX_ENTRIES:
            raise ValueError('Bilgi kütüphanesi doldu')
        content=(json.dumps(data,ensure_ascii=False,indent=2)+'\n').encode('utf-8')
        if len(content)>MAX_STORE_BYTES: raise ValueError('Bilgi kütüphanesi dosyası çok büyük')
        # Lock file denies concurrent writers instead of silently losing updates.
        lock=self.path.with_name(self.path.name+'.lock')
        try: fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
        except FileExistsError: raise RuntimeError('Başka bir düzenleme sürüyor; tekrar deneyin') from None
        os.close(fd)
        tmp=None
        try:
            current=self.read()
            # Caller must work from the last read revision (optimistic concurrency).
            if data.get('_previous_revision') != current.get('_revision'):
                raise RuntimeError('Başka pencerede kütüphane değişti; yeniden açın')
            data.pop('_previous_revision',None)
            data['_revision']=uuid4().hex
            content=(json.dumps(data,ensure_ascii=False,indent=2)+'\n').encode('utf-8')
            if len(content)>MAX_STORE_BYTES: raise ValueError('Kütüphane boyut sınırı aşıldı')
            tfd,tmp=tempfile.mkstemp(prefix='.coach_kb_',dir=self.path.parent)
            with os.fdopen(tfd,'wb') as f:
                f.write(content);f.flush();os.fsync(f.fileno())
            os.replace(tmp,self.path);tmp=None
        finally:
            if tmp: Path(tmp).unlink(missing_ok=True)
            lock.unlink(missing_ok=True)

    def import_document(self,source):
        source=Path(source)
        fragments=_extract_document(source)
        if not fragments: raise ValueError('İncelemeye uygun paragraf bulunamadı')
        data=self.read(); data['_previous_revision']=data.get('_revision')
        if len(data['entries'])+len(fragments)>MAX_ENTRIES:raise ValueError('Kütüphane kayıt sınırı aşıldı')
        sha=hashlib.sha256(source.read_bytes()).hexdigest()
        for item in fragments:
            data['entries'].append({'id':uuid4().hex,'status':'draft','source_name':source.name[:180],
               'source_sha256':sha,'source_type':'teacher_document','imported_at':_now(),
               'original_excerpt':item,'title':'','category':'coaching','keywords':'',
               'guidance':'','teacher_question':'','caution':'',
               'automatic_action_allowed':False,'sensitive':False})
        self._save(data)
        return len(fragments)

    def import_official_snapshot(self, snapshot):
        """Save one official-page excerpt as a DRAFT; never activate automatically.

        A website snapshot is untrusted informational material, not instructions,
        medical/psychological evidence, or a student measurement.
        """
        from services.coach_assistant_v1.official_sources_v127 import catalog_entry
        if not isinstance(snapshot,dict):raise ValueError('Kaynak kaydı geçersiz')
        info=catalog_entry(snapshot.get('source_id'))
        if snapshot.get('url')!=info['url'] or snapshot.get('automatic_action_allowed') is not False:
            raise ValueError('Kaynak bağlantısı doğrulanamadı')
        stamp=_safe_str(snapshot.get('checked_at',''),60)
        if not stamp or not re.match(r'^\d{4}-\d{2}-\d{2}T',stamp):
            raise ValueError('Kaynak kontrol tarihi geçersiz')
        digest=str(snapshot.get('sha256',''))
        if not re.fullmatch(r'[0-9a-f]{64}',digest):raise ValueError('Kaynak izi geçersiz')
        excerpt=_safe_str(snapshot.get('excerpt',''),800)
        if len(excerpt)<20:raise ValueError('Kaynak açıklaması çok kısa')
        data=self.read();data['_previous_revision']=data.get('_revision')
        if len(data['entries'])>=MAX_ENTRIES:raise ValueError('Kütüphane kayıt sınırı aşıldı')
        # Duplicate imports of the SAME checked version are skipped.
        for e in data['entries']:
            if (e.get('source_type')=='official_web_snapshot'
                and e.get('source_url')==info['url'] and e.get('source_sha256')==digest):
                return False
        data['entries'].append({
            'id':uuid4().hex,'status':'draft','source_name':info['name'],
            'source_url':info['url'],'source_checked_at':stamp,
            'source_page_date':snapshot.get('declared_publish_date') or '',
            'source_sha256':digest,'source_type':'official_web_snapshot',
            'imported_at':_now(),'original_excerpt':excerpt,
            'title':'','category':info['category'],'keywords':'',
            'guidance':'','teacher_question':'','caution':'',
            'automatic_action_allowed':False,'sensitive':False,
        })
        self._save(data)
        return True

    def list_entries(self, status=None):
        entries=self.read()['entries']
        return [e.copy() for e in entries if status is None or e['status']==status]

    def review(self,entry_id,*,title,category,keywords,guidance,teacher_question,caution=''):
        title=_safe_str(title,150);guidance=_safe_str(guidance,700)
        teacher_question=_safe_str(teacher_question,250);caution=_safe_str(caution,350)
        keywords=_safe_str(keywords,250)
        if not all((title, guidance, teacher_question, keywords)):
            raise ValueError('Başlık, öneri, soru ve anahtar kelimeler zorunlu')
        if category not in CATEGORIES: raise ValueError('Bilinmeyen kategori')
        data=self.read();data['_previous_revision']=data.get('_revision')
        e=next((x for x in data['entries'] if x['id']==entry_id),None)
        if e is None or e['status'] != 'draft':raise ValueError('Yalnızca taslak kayıt onaylanabilir')
        e.update(title=title,category=category,keywords=keywords,guidance=guidance,
                 teacher_question=teacher_question,caution=caution,status='reviewed',
                 reviewed_at=_now(),review_type='teacher_editorial_review',
                 scientific_validation='not_claimed',automatic_action_allowed=False)
        self._save(data)

    def discard_draft(self,entry_id):
        data=self.read();data['_previous_revision']=data.get('_revision')
        e=next((x for x in data['entries'] if x['id']==entry_id),None)
        if e is None or e['status']!='draft':
            raise ValueError('Yalnızca inceleme taslağı silinebilir')
        data['entries']=[item for item in data['entries'] if item['id']!=entry_id]
        self._save(data)

    def retire(self,entry_id):
        data=self.read();data['_previous_revision']=data.get('_revision')
        e=next((x for x in data['entries'] if x['id']==entry_id),None)
        if e is None or e['status']!='reviewed':raise ValueError('Yalnızca etkin bilgi pasifleştirilebilir')
        e['status']='retired';e['retired_at']=_now()
        self._save(data)

    def search(self,query,limit=1):
        tokens=set(fold(query[:600]).split())
        if len(' '.join(tokens))<4:return ()
        if any(t.startswith(prefix) for t in tokens for prefix in ('kaygi','depresyon','panik','travma','uyku','stres')):
            return ()
        matches=[]
        for e in self.list_entries('reviewed'):
            if e.get('sensitive') or e.get('automatic_action_allowed') is not False:continue
            if e.get('category') not in CATEGORIES:continue
            try:
                title=_safe_str(e['title'],150)
                guidance=_safe_str(e['guidance'],700)
                teacher_question=_safe_str(e['teacher_question'],250)
                keywords=_safe_str(e['keywords'],250)
            except (ValueError, KeyError):continue
            if not title or not guidance or not teacher_question:continue
            kw=[fold(x.strip()) for x in keywords.split(',') if x.strip()]
            score=sum(8 for k in kw if k and (k in fold(query) and len(k)>=4))
            score+=2*len(tokens & set(fold(title).split()))
            if score>=8:matches.append((score,e))
        matches.sort(key=lambda pair:(-pair[0],pair[1]['id']))
        return tuple(entry for _,entry in matches[:max(0,min(int(limit),2))])


def enrich_with_reviewed(answer,question,store=None):
    """Add at most one non-executable teacher-approved general reference."""
    from services.coach_assistant_v1.dialogue import CoachAnswer
    if not isinstance(answer,CoachAnswer):raise TypeError('Geçersiz Koç Asistanı yanıtı')
    store=store or ReviewStore()
    matches=store.search(question,limit=1)
    if not matches:return answer
    x=matches[0]
    note='\n\n📚 Kurum/öğretmen tarafından incelenmiş rehber (bilimsel doğrulama veya öğrenci ölçümü değildir): '+x['guidance']
    if x.get('teacher_question'):note+='\nGörüşme sorusu: '+x['teacher_question']
    if x.get('caution'):note+='\nSınır: '+x['caution']
    source_info='Öğretmen belge referansı: '+str(x['source_name'])[:180]+' · öğretmen incelemesi'
    if x.get('source_type')=='official_web_snapshot':
        from services.coach_assistant_v1.official_sources_v127 import SOURCES
        if x.get('source_url') in [entry['url'] for entry in SOURCES.values()]:
            source_info+=' · '+x['source_url']+' · kontrol: '+str(x.get('source_checked_at',''))[:10]
    evidence=(*answer.evidence,source_info)
    return CoachAnswer(answer.title,answer.message+note,evidence,answer.action_target)
