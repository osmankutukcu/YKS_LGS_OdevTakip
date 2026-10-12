"""V20.1: transparent, deterministic card-layout detection (read-only).

No publisher identity inference, no OCR, no student/exam guesses, no writes.
Unknown or ambiguous headings are handled conservatively by callers.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import zipfile

from services.trial_exam_import_v17 import _key, MAX_FIELD_LENGTH, MAX_FILE_BYTES, MAX_INPUT_ROWS

# Exact aliases only. Never try fuzzy matches for different curriculum topics.
_ALIASES = {
    'ogrenci_id': ('ogrenci id', 'ogrenci_id', 'ogrenci no', 'ogrenci numarasi', 'okul no', 'student id'),
    'deneme_id': ('deneme id', 'deneme_id', 'sinav id', 'sinav no', 'exam id'),
    'ders': ('ders', 'ders adi', 'ders_adi', 'brans', 'branş', 'subject', 'test adi', 'test'),
    'soru_no': ('soru no', 'soru_no', 'soru numarasi', 'soru num', 'soru', 'question no', 'question', 'soru #'),
    'sonuc': ('sonuc', 'cevap durumu', 'dogru yanlis bos', 'soru sonucu', 'durum', 'result',
              'isaret durumu', 'soru durumu', 'd/y/b', 'd y b'),
    'konu': ('konu', 'konu adi', 'konu_adi', 'topic', 'konu basligi', 'alt konu'),
    'kazanim_kodu': ('kazanim', 'kazanim kodu', 'kazanim_kodu', 'ogrenme ciktisi',
                     'outcome', 'kazanim no', 'kazanim id'),
}


def _normalize(text: object) -> str:
    return _key(str(text or '')).replace('_', ' ').strip().strip(':').strip()


_INDEX = {}
for _field, _names in _ALIASES.items():
    for _name in _names:
        _alias = _normalize(_name)
        if _alias in _INDEX and _INDEX[_alias] != _field:
            raise RuntimeError('Conflicting built-in column aliases: ' + _alias)
        _INDEX[_alias] = _field


@dataclass(frozen=True)
class HeaderMatch:
    columns: dict[str, int]
    layout: str
    certainty: str  # structural only, never content accuracy


def match_header(row: list[object] | tuple[object, ...]) -> HeaderMatch | None:
    """Match strict column labels; reject duplicate synonym columns.

    A column called 'Net' alone is never sufficient evidence for a question row.
    """
    mapping = {}
    for idx, cell in enumerate(row):
        norm = _normalize(cell)
        field = _INDEX.get(norm)
        if field is None:
            continue
        if field in mapping:
            raise ValueError(f'Tekrarlanan/çakışan sütun: {field}')
        mapping[field] = idx
    if not all(k in mapping for k in ('soru_no', 'sonuc', 'konu')):
        return None
    label = 'Dersli soru tablosu' if 'ders' in mapping else 'Ders başlığına bağlı soru tablosu'
    return HeaderMatch(mapping, label, 'sütun eşleştirmesi; satırların doğruluğu ayrıca kontrol edilmeli')


def read_xlsx_sheets(path: str | Path) -> tuple[tuple[str, list[list[object]]], ...]:
    """Read visible XLSX sheets only; never execute macros/formulas/external links.

    Scanning all sheets fixes the previous first-sheet-only limitation. Hard limits
    protect against huge workbooks. A formula anywhere in an inspected visible
    sheet rejects the workbook; it is not treated as an authoritative score.
    """
    path = Path(path)
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_FILE_BYTES:
        raise ValueError('Excel dosyası boş veya 8 MB sınırını aşıyor.')
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(i.file_size for i in entries) > 40 * 1024 * 1024:
            raise ValueError('Excel arşivi güvenlik sınırlarını aşıyor.')
        if any('vbaproject.bin' in n.lower() or n.lower().endswith('.bin') for n in archive.namelist()):
            raise ValueError('Makro/ikili ek içeren Excel desteklenmiyor.')
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Çok sayfalı Excel için openpyxl kurulmalıdır.') from exc
    wb = load_workbook(filename=path, read_only=True, data_only=False, keep_links=False)
    try:
        visible = [sheet for sheet in wb.worksheets if sheet.sheet_state == 'visible']
        if not visible or len(visible) > 15:
            raise ValueError('1–15 görünür çalışma sayfası desteklenir.')
        total_rows = 0
        result = []
        for sheet in visible:
            grid = []
            for cells in sheet.iter_rows():
                total_rows += 1
                if total_rows > MAX_INPUT_ROWS + 150:
                    raise ValueError('Excel çalışma kitabında çok fazla satır var.')
                if len(cells) > 64:
                    raise ValueError('Excel çalışma sayfasında en fazla 64 sütun desteklenir.')
                values = []
                for cell in cells:
                    if cell.data_type == 'f':
                        raise ValueError('Formül içeren Excel hücresi desteklenmiyor; değer olarak yapıştırın.')
                    value = '' if cell.value is None else str(cell.value)
                    if len(value) > MAX_FIELD_LENGTH:
                        raise ValueError('Excel hücresi 500 karakter sınırını aşıyor.')
                    values.append(value)
                grid.append(values)
            result.append((sheet.title, grid))
        return tuple(result)
    finally:
        wb.close()


def has_partial_question_header(grid: list[list[object]]) -> bool:
    """Detect broken question tables; never silently skip a potentially useful sheet."""
    for row in grid[:50]:
        fields = {_INDEX[x] for cell in row if (x := _normalize(cell)) in _INDEX}
        if 'soru_no' in fields and ('sonuc' in fields or 'konu' in fields):
            return True
    return False
